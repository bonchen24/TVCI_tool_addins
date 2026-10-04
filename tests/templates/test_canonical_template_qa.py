import importlib.util
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.shared import Mm


ROOT = Path(__file__).resolve().parents[2]
GENERATOR_PATH = ROOT / "scripts" / "generate-canonical-templates.py"
QA_PATH = ROOT / "scripts" / "verify-canonical-templates.py"
SPEC_PATH = ROOT / "templates" / "specs" / "administrative-templates.v1.json"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise AssertionError(f"Cannot load {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class CanonicalTemplateQaTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.generator = load_module("canonical_generator_for_qa_tests", GENERATOR_PATH)
        cls.qa = load_module("canonical_qa_for_tests", QA_PATH) if QA_PATH.exists() else None
        cls.spec = cls.generator.load_spec(SPEC_PATH)
        cls.templates = {template["key"]: template for template in cls.spec["templates"]}
        cls.temporary = tempfile.TemporaryDirectory()
        cls.output_dir = Path(cls.temporary.name) / "generated"
        cls.generator.generate_canonical_templates(SPEC_PATH, cls.output_dir)

    @classmethod
    def tearDownClass(cls):
        cls.temporary.cleanup()

    def require_qa(self):
        self.assertIsNotNone(self.qa, "verify-canonical-templates.py must be implemented after this red test")
        return self.qa

    def canonical_path(self, key: str) -> Path:
        return self.output_dir / self.templates[key]["canonicalFile"]

    def assert_finding(self, path: Path, key: str, code: str):
        report = self.require_qa().verify_docx(path, self.spec, self.templates[key])
        self.assertIn(code, {finding["code"] for finding in report["findings"]})

    def test_all_23_generated_packages_pass_structural_and_semantic_qa(self):
        qa = self.require_qa()
        results = [qa.verify_docx(self.canonical_path(key), self.spec, template) for key, template in self.templates.items()]
        failures = {result["templateKey"]: result["findings"] for result in results if result["findings"]}
        self.assertEqual(failures, {})
        self.assertEqual(len(results), 23)
        self.assertTrue(all(result["structuralQa"]["status"] == "passed" for result in results))
        self.assertTrue(all(result["semanticQa"]["status"] == "passed" for result in results))

    def test_rejects_a_document_with_wrong_page_margins(self):
        path = self.canonical_path("iemm-cong-van")
        broken = Document(path)
        broken.sections[0].left_margin = Mm(50)
        target = Path(self.temporary.name) / "wrong-margins.docx"
        broken.save(target)
        self.assert_finding(target, "iemm-cong-van", "page-margin")

    def test_rejects_a_two_column_header_that_exceeds_the_text_width(self):
        path = self.canonical_path("iemm-cong-van")
        broken = Document(path)
        broken.tables[0].columns[0].width = Mm(100)
        broken.tables[0].columns[1].width = Mm(100)
        target = Path(self.temporary.name) / "wide-header.docx"
        broken.save(target)
        self.assert_finding(target, "iemm-cong-van", "header-width")

    def test_rejects_a_page_number_centered_on_the_asymmetric_text_area(self):
        path = self.canonical_path("iemm-cong-van")
        broken = Document(path)
        paragraph = broken.sections[0].header.paragraphs[0]
        paragraph.paragraph_format.left_indent = Mm(0)
        paragraph.paragraph_format.right_indent = Mm(0)
        target = Path(self.temporary.name) / "uncentered-page-number.docx"
        broken.save(target)
        self.assert_finding(target, "iemm-cong-van", "page-number-position")

    def test_rejects_a_categorized_document_title_on_a_letter(self):
        path = self.canonical_path("iemm-cong-van")
        broken = Document(path)
        broken.add_paragraph("CÔNG VĂN")
        target = Path(self.temporary.name) / "letter-with-title.docx"
        broken.save(target)
        self.assert_finding(target, "iemm-cong-van", "forbidden-label")

    def test_rejects_a_decision_missing_its_articles(self):
        path = self.canonical_path("iemm-quyet-dinh-ca-biet")
        broken = Document(path)
        for paragraph in broken.paragraphs:
            if "Điều 2." in paragraph.text:
                paragraph.text = "Khoản đã bị thay đổi"
        target = Path(self.temporary.name) / "decision-without-article.docx"
        broken.save(target)
        self.assert_finding(target, "iemm-quyet-dinh-ca-biet", "missing-label")

    def test_rejects_a_proposal_missing_its_recipient(self):
        path = self.canonical_path("iemm-to-trinh-vien")
        broken = Document(path)
        for paragraph in broken.paragraphs:
            if "Kính gửi:" in paragraph.text:
                paragraph.text = "Nội dung đã thay đổi"
        target = Path(self.temporary.name) / "proposal-without-recipient.docx"
        broken.save(target)
        self.assert_finding(target, "iemm-to-trinh-vien", "missing-label")

    def test_rejects_a_copy_without_its_certification_block(self):
        path = self.canonical_path("iemm-ban-sao")
        broken = Document(path)
        for paragraph in broken.paragraphs:
            if "SAO Y BẢN CHÍNH" in paragraph.text:
                paragraph.text = "Nội dung chứng thực bị thiếu"
        target = Path(self.temporary.name) / "copy-without-certification.docx"
        broken.save(target)
        self.assert_finding(target, "iemm-ban-sao", "missing-label")

    def test_rejects_duplicate_content_control_tags(self):
        source = self.canonical_path("iemm-cong-van")
        target = Path(self.temporary.name) / "duplicate-tags.docx"
        with zipfile.ZipFile(source, "r") as original, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output:
            for item in original.infolist():
                data = original.read(item.filename)
                if item.filename == "word/document.xml":
                    root = ET.fromstring(data)
                    body = root.find(f"{{{W_NS}}}body")
                    for _ in range(2):
                        control = ET.SubElement(body, f"{{{W_NS}}}sdt")
                        properties = ET.SubElement(control, f"{{{W_NS}}}sdtPr")
                        tag = ET.SubElement(properties, f"{{{W_NS}}}tag")
                        tag.set(f"{{{W_NS}}}val", "DUPLICATE_TAG")
                        ET.SubElement(control, f"{{{W_NS}}}sdtContent")
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                output.writestr(item.filename, data)
        self.assert_finding(target, "iemm-cong-van", "duplicate-control-tag")


if __name__ == "__main__":
    unittest.main()
