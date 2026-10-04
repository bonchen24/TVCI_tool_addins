#!/usr/bin/env python3
"""Generate deterministic canonical administrative DOCX files from the versioned JSON spec."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import platform
import shutil
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import docx
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


GENERATOR_VERSION = "1.0.0"
ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SPEC = ROOT / "templates" / "specs" / "administrative-templates.v1.json"
DEFAULT_OUTPUT = ROOT / "canonical_templates" / "generated"
FIXED_PROPERTY_DATE = datetime(2000, 1, 1, tzinfo=timezone.utc)
ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    return sha256_bytes(path.read_bytes())


def load_spec(path: Path = DEFAULT_SPEC) -> dict[str, Any]:
    spec = json.loads(path.read_text(encoding="utf-8"))
    if spec.get("generatorVersion") != GENERATOR_VERSION:
        raise ValueError(
            f"Spec generatorVersion {spec.get('generatorVersion')!r} does not match {GENERATOR_VERSION}."
        )
    if not spec.get("ruleSpecVersion") or not spec.get("templates") or not spec.get("documentTypes"):
        raise ValueError("The template spec must define versioned rules, document types, and templates.")
    return spec


def _mm_to_twips(value: float) -> int:
    return int(round(value * 1440 / 25.4))


def _set_run_font(run, style: dict[str, Any], common: dict[str, Any]) -> None:
    font_name = common["font"]["name"]
    run.font.name = font_name
    run.font.size = Pt(style.get("fontPt", common["font"]["bodyPt"]))
    run.font.bold = bool(style.get("bold", False))
    run.font.italic = bool(style.get("italic", False))
    run.font.color.rgb = RGBColor.from_string(style.get("color", common["font"]["color"]))
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        r_fonts.set(qn(f"w:{name}"), font_name)
    language = r_pr.find(qn("w:lang"))
    if language is None:
        language = OxmlElement("w:lang")
        r_pr.append(language)
    language.set(qn("w:val"), "vi-VN")


def _set_bottom_rule(paragraph, common: dict[str, Any]) -> None:
    rule = common["horizontalRule"]
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), rule["style"])
    bottom.set(qn("w:sz"), str(rule["sizeEighthPt"]))
    bottom.set(qn("w:space"), str(rule["spacePt"]))
    bottom.set(qn("w:color"), rule["color"])
    borders.append(bottom)


def _format_paragraph(paragraph, style: dict[str, Any], common: dict[str, Any]) -> None:
    paragraph.alignment = ALIGNMENTS[style.get("alignment", "left")]
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(style.get("spaceBeforePt", 0))
    fmt.space_after = Pt(style.get("spaceAfterPt", 0))
    if "firstLineIndentMm" in style:
        fmt.first_line_indent = Mm(style["firstLineIndentMm"])
    if "lineSpacingMultiple" in style:
        fmt.line_spacing = style["lineSpacingMultiple"]
    if style.get("keepWithNext"):
        fmt.keep_with_next = True
    if style.get("bottomRule"):
        _set_bottom_rule(paragraph, common)


def _add_text(paragraph, text: str, style: dict[str, Any], common: dict[str, Any]) -> None:
    rendered_text = text.upper() if style.get("uppercase") else text
    run = paragraph.add_run(rendered_text)
    _set_run_font(run, style, common)


def _add_paragraph(container, text: str, style: dict[str, Any], common: dict[str, Any]):
    paragraph = container.add_paragraph()
    _format_paragraph(paragraph, style, common)
    _add_text(paragraph, text, style, common)
    return paragraph


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def _set_cell_margins(cell, values: dict[str, float]) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side in ("top", "start", "bottom", "end"):
        item = margins.find(qn(f"w:{side}"))
        if item is None:
            item = OxmlElement(f"w:{side}")
            margins.append(item)
        item.set(qn("w:w"), str(_mm_to_twips(values[side])))
        item.set(qn("w:type"), "dxa")


def _layout_table(container, common: dict[str, Any], columns: int = 2):
    table = container.add_table(rows=1, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_table_borders(table)
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        _set_cell_margins(cell, common["tableCellMarginsMm"])
    return table


def _set_table_widths(table, widths_mm: list[float]) -> None:
    for index, width in enumerate(widths_mm):
        table.columns[index].width = Mm(width)
        for cell in table.columns[index].cells:
            cell.width = Mm(width)


def _cell_paragraph(cell, text: str, style: dict[str, Any], common: dict[str, Any], first: bool = False):
    paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
    if first:
        paragraph.clear()
    _format_paragraph(paragraph, style, common)
    if text:
        _add_text(paragraph, text, style, common)
    return paragraph


def _page_field(section, style: dict[str, Any], common: dict[str, Any]) -> None:
    paragraph = section.header.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = ALIGNMENTS[style["position"].replace("top-", "")]
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    center_compensation_mm = (section.right_margin.mm - section.left_margin.mm) / 2
    paragraph.paragraph_format.left_indent = Mm(center_compensation_mm)
    paragraph.paragraph_format.right_indent = Mm(-center_compensation_mm)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    font_size = OxmlElement("w:sz")
    font_size.set(qn("w:val"), str(int(style["fontPt"] * 2)))
    properties.append(font_size)
    fonts = OxmlElement("w:rFonts")
    font_name = common["font"]["name"]
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{name}"), font_name)
    properties.append(fonts)
    run.append(properties)
    value = OxmlElement("w:t")
    value.text = "2"
    run.append(value)
    field.append(run)
    paragraph._p.append(field)


def _add_document_header(document, spec: dict[str, Any], template: dict[str, Any]) -> None:
    common = spec["commonStyle"]
    organization = spec["organizations"][template["organization"]]
    section = document.sections[0]
    page = common["page"]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(page["widthMm"])
    section.page_height = Mm(page["heightMm"])
    margins = page["marginsMm"]
    section.top_margin = Mm(margins["top"])
    section.bottom_margin = Mm(margins["bottom"])
    section.left_margin = Mm(margins["left"])
    section.right_margin = Mm(margins["right"])
    section.header_distance = Mm(common["headerFooter"]["headerDistanceMm"])
    section.footer_distance = Mm(common["headerFooter"]["footerDistanceMm"])
    section.different_first_page_header_footer = bool(common["pageNumber"]["suppressFirstPage"])
    section.first_page_header.paragraphs[0].clear()
    if not section.different_first_page_header_footer:
        _page_field(section, common["pageNumber"], common)
    else:
        _page_field(section, common["pageNumber"], common)

    usable_width = page["widthMm"] - margins["left"] - margins["right"]
    grid = common["headerGrid"]
    table = _layout_table(document, common, columns=grid["columns"])
    left_width = round(usable_width * grid["leftWidthRatio"], 2)
    right_width = round(usable_width - left_width, 2)
    _set_table_widths(table, [left_width, right_width])

    left = table.cell(0, 0)
    _cell_paragraph(left, organization["parentAgency"], common["organizationParent"], common, first=True)
    _cell_paragraph(left, organization["issuingAgency"], common["organizationIssuer"], common)
    if organization.get("department"):
        _cell_paragraph(left, organization["department"], common["organizationDepartment"], common)

    if spec["documentTypes"][template["typeSchema"]].get("headerMode", "administrative") != "internal":
        right = table.cell(0, 1)
        _cell_paragraph(right, common["nationalHeading"]["text"], common["nationalHeading"], common, first=True)
        _cell_paragraph(right, common["nationalMotto"]["text"], common["nationalMotto"], common)


def _add_number_date(document, spec: dict[str, Any], template: dict[str, Any], block: dict[str, Any]) -> None:
    common = spec["commonStyle"]
    page = common["page"]
    margins = page["marginsMm"]
    usable_width = page["widthMm"] - margins["left"] - margins["right"]
    table = _layout_table(document, common, columns=2)
    _set_table_widths(table, [usable_width / 2, usable_width / 2])
    symbol_pattern = block.get("symbolPattern", template.get("symbolPattern"))
    if not symbol_pattern:
        raise ValueError(f"{template['key']} needs a symbolPattern or an explicit internal form schema.")
    _cell_paragraph(table.cell(0, 0), symbol_pattern, common["symbol"], common, first=True)
    place_date = f"{{DIA_DANH}}, ngày {{NGAY_THANG}}"
    _cell_paragraph(table.cell(0, 1), place_date, common["placeDate"], common, first=True)


def _add_signature_cell(cell, role: str, name: str, spec: dict[str, Any], first: bool = True) -> None:
    common = spec["commonStyle"]
    _cell_paragraph(cell, role, common["signatureRole"], common, first=first)
    _cell_paragraph(cell, "", common["signatureRole"], common)
    _cell_paragraph(cell, name, common["signatureName"], common)


def _add_recipients_signature(document, spec: dict[str, Any]) -> None:
    common = spec["commonStyle"]
    page = common["page"]
    margins = page["marginsMm"]
    width = page["widthMm"] - margins["left"] - margins["right"]
    table = _layout_table(document, common, columns=2)
    _set_table_widths(table, [width / 2, width / 2])
    left = table.cell(0, 0)
    _cell_paragraph(left, "Nơi nhận:", common["recipientHeading"], common, first=True)
    _cell_paragraph(left, "- {NOI_NHAN}", common["recipientBody"], common)
    right = table.cell(0, 1)
    _add_signature_cell(right, "{CHUC_VU_NGUOI_KY}", "{HO_TEN_NGUOI_KY}", spec)


def _add_single_signature(document, spec: dict[str, Any]) -> None:
    common = spec["commonStyle"]
    page = common["page"]
    margins = page["marginsMm"]
    width = page["widthMm"] - margins["left"] - margins["right"]
    table = _layout_table(document, common, columns=2)
    _set_table_widths(table, [width / 2, width / 2])
    _add_signature_cell(table.cell(0, 1), "{CHUC_VU_NGUOI_KY}", "{HO_TEN_NGUOI_KY}", spec)


def _add_named_signatures(document, spec: dict[str, Any], roles: list[str]) -> None:
    common = spec["commonStyle"]
    page = common["page"]
    margins = page["marginsMm"]
    width = page["widthMm"] - margins["left"] - margins["right"]
    table = _layout_table(document, common, columns=len(roles))
    _set_table_widths(table, [width / len(roles)] * len(roles))
    for cell, role in zip(table.rows[0].cells, roles):
        _add_signature_cell(cell, role, "{HO_TEN_NGUOI_KY}", spec)


def _add_copy_certification(document, spec: dict[str, Any]) -> None:
    common = spec["commonStyle"]
    style = common["copyCertification"]
    _add_paragraph(document, "SAO Y BẢN CHÍNH", style, common)
    _add_paragraph(document, "Số: {SO_VAN_BAN}     Ngày: {NGAY_THANG}", common["recipientLine"], common)
    _add_paragraph(document, "{CHUC_VU_NGUOI_KY}", common["signatureRole"], common)
    _add_paragraph(document, "{HO_TEN_NGUOI_KY}", common["signatureName"], common)


def _add_block(document, spec: dict[str, Any], template: dict[str, Any], block: dict[str, Any]) -> None:
    common = spec["commonStyle"]
    kind = block["kind"]
    if kind == "number_date":
        _add_number_date(document, spec, template, block)
    elif kind == "place_date":
        style = common["placeDate"]
        _add_paragraph(document, "{DIA_DANH}, ngày {NGAY_THANG}", style, common)
    elif kind in ("title", "subject", "paragraph", "label"):
        style_name = block.get("style") or {
            "title": "documentTitle",
            "subject": "subject",
            "paragraph": "bodyParagraph",
            "label": "bodyLabel",
        }[kind]
        style = common[style_name]
        text = block.get("text", "")
        paragraph = document.add_paragraph()
        _format_paragraph(paragraph, style, common)
        _add_text(paragraph, text, style, common)
    elif kind == "recipient":
        _add_paragraph(document, "Kính gửi: {NOI_NHAN}", common["recipientLine"], common)
    elif kind == "recipients_signature":
        _add_recipients_signature(document, spec)
    elif kind == "signature":
        _add_single_signature(document, spec)
    elif kind == "meeting_signatures":
        _add_named_signatures(document, spec, ["CHỦ TRÌ", "THƯ KÝ"])
    elif kind == "leave_signatures":
        _add_named_signatures(document, spec, ["NGƯỜI LÀM ĐƠN", "TRƯỞNG ĐƠN VỊ", "LÃNH ĐẠO PHÊ DUYỆT"])
    elif kind == "commitment_signatures":
        _add_named_signatures(document, spec, ["NGƯỜI CAM KẾT", "XÁC NHẬN CỦA ĐƠN VỊ"])
    elif kind == "copy_certification":
        _add_copy_certification(document, spec)
    else:
        raise ValueError(f"Unsupported block kind {kind!r} in {template['key']}.")


def build_document(spec: dict[str, Any], template: dict[str, Any]) -> Document:
    document = Document()
    common = spec["commonStyle"]
    normal = document.styles["Normal"]
    normal.font.name = common["font"]["name"]
    normal.font.size = Pt(common["font"]["bodyPt"])
    normal.font.color.rgb = RGBColor.from_string(common["font"]["color"])
    normal.paragraph_format.space_before = Pt(common["body"]["spaceBeforePt"])
    normal.paragraph_format.space_after = Pt(common["body"]["spaceAfterPt"])
    normal.paragraph_format.line_spacing = common["body"]["lineSpacingMultiple"]
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), common["font"]["name"])

    type_spec = spec["documentTypes"][template["typeSchema"]]
    _add_document_header(document, spec, template)
    for block in type_spec["blocks"]:
        _add_block(document, spec, template, block)

    properties = document.core_properties
    properties.title = template["name"]
    properties.subject = template["name"]
    properties.author = "TVCI Word Add-in canonical generator"
    properties.last_modified_by = "TVCI Word Add-in canonical generator"
    properties.keywords = "generated canonical; administrative template"
    properties.comments = f"ruleSpecVersion={spec['ruleSpecVersion']}; generatorVersion={GENERATOR_VERSION}"
    properties.identifier = template["key"]
    properties.revision = 1
    properties.created = FIXED_PROPERTY_DATE
    properties.modified = FIXED_PROPERTY_DATE
    return document


def normalize_docx(package_bytes: bytes) -> bytes:
    source = zipfile.ZipFile(io.BytesIO(package_bytes), "r")
    output = io.BytesIO()
    with source, zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target:
        target.comment = b""
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.create_system = 0
            info.external_attr = 0
            info.flag_bits = 0
            info.extra = b""
            info.comment = b""
            target.writestr(info, source.read(name), compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
    return output.getvalue()


def _source_evidence(spec: dict[str, Any], source_keys: list[str]) -> list[dict[str, Any]]:
    sources = {source["key"]: source for source in spec["sources"]}
    return [
        {
            "id": key,
            "label": sources[key]["label"],
            "reference": sources[key].get("reference"),
            "driveFileId": sources[key].get("driveFileId"),
            "parentDriveFileId": sources[key].get("parentDriveFileId"),
            "priority": sources[key]["priority"],
        }
        for key in source_keys
    ]


def _template_manifest_record(spec: dict[str, Any], template: dict[str, Any], canonical_hash: str) -> dict[str, Any]:
    return {
        "key": template["key"],
        "name": template["name"],
        "organization": template["organization"],
        "typeSchema": template["typeSchema"],
        "canonicalFile": template["canonicalFile"],
        "canonicalPath": f"canonical_templates/generated/{template['canonicalFile']}",
        "canonicalSha256": canonical_hash,
        "runtimePath": template["runtimePath"],
        "runtimeSha256": None,
        "comparison": None,
        "catalogIds": template["catalogIds"],
        "normativeSources": _source_evidence(spec, template["normativeSources"]),
        "referenceSources": _source_evidence(spec, template["referenceSources"]),
        "sourceClassification": template.get("sourceClassification", "rule-derived"),
        "sourceNotes": template["sourceNotes"],
        "structuralQa": {"status": "unverified", "reportPath": None, "sha256": None},
        "semanticQa": {"status": "unverified", "reportPath": None, "sha256": None},
        "deterministicQa": {"status": "unverified", "reportPath": None, "sha256": None},
        "visualQa": {"status": "unverified", "renderer": None, "reportPath": None, "sha256": None},
        "runtimeParity": {"status": "unverified", "comparison": None, "runtimeSha256": None},
    }


def generate_canonical_templates(spec_path: Path = DEFAULT_SPEC, output_dir: Path = DEFAULT_OUTPUT) -> dict[str, Any]:
    spec = load_spec(spec_path)
    script_hash = sha256_file(Path(__file__).resolve())
    output_dir.mkdir(parents=True, exist_ok=True)
    expected_files = {template["canonicalFile"] for template in spec["templates"]}
    existing_docx = {path.relative_to(output_dir).as_posix() for path in output_dir.rglob("*.docx")}
    unexpected = existing_docx - expected_files
    if unexpected:
        raise ValueError(f"Refusing to leave unexpected DOCX files in generated output: {sorted(unexpected)}")

    records = []
    for template in spec["templates"]:
        destination = output_dir / Path(template["canonicalFile"])
        destination.parent.mkdir(parents=True, exist_ok=True)
        package = io.BytesIO()
        build_document(spec, template).save(package)
        normalized = normalize_docx(package.getvalue())
        destination.write_bytes(normalized)
        records.append(_template_manifest_record(spec, template, sha256_bytes(normalized)))

    manifest = {
        "formatVersion": 1,
        "ruleSpecVersion": spec["ruleSpecVersion"],
        "generatorVersion": GENERATOR_VERSION,
        "generatorSha256": script_hash,
        "toolchain": {
            "pythonVersion": platform.python_version(),
            "pythonDocxVersion": docx.__version__,
            "zlibVersion": zipfile.zlib.ZLIB_VERSION,
        },
        "sourceInventory": _source_evidence(spec, [source["key"] for source in spec["sources"]]),
        "templates": records,
    }
    (output_dir / "provenance.json").write_text(
        json.dumps(manifest, ensure_ascii=False, sort_keys=True, indent=2) + "\n",
        encoding="utf-8",
        newline="\n",
    )
    return manifest


def _read_json(path: Path) -> dict[str, Any]:
    value = json.loads(Path(path).read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise ValueError(f"Expected a JSON object in {path}.")
    return value


def _write_json(path: Path, value: dict[str, Any]) -> None:
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    Path(path).write_text(json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2) + "\n", encoding="utf-8", newline="\n")


def _report_path(path: Path) -> str:
    try:
        return Path(path).resolve().relative_to(ROOT.resolve()).as_posix()
    except ValueError:
        return str(Path(path).resolve())


def _report_index(report: dict[str, Any], report_name: str, expected_keys: set[str]) -> dict[str, dict[str, Any]]:
    if report.get("status") != "passed":
        raise ValueError(f"{report_name} QA report has not passed.")
    records = report.get("templates")
    if not isinstance(records, list):
        raise ValueError(f"{report_name} QA report has no template records.")
    indexed = {item.get("templateKey"): item for item in records if isinstance(item, dict)}
    if len(indexed) != len(records) or set(indexed) != expected_keys:
        raise ValueError(f"{report_name} QA report must contain every spec template exactly once.")
    return indexed


def verify_determinism(spec_path: Path = DEFAULT_SPEC, input_dir: Path = DEFAULT_OUTPUT) -> dict[str, Any]:
    spec = load_spec(spec_path)
    input_dir = Path(input_dir)
    manifest_path = input_dir / "provenance.json"
    if not manifest_path.is_file():
        raise ValueError(f"Canonical provenance manifest is missing: {manifest_path}")
    manifest = _read_json(manifest_path)
    expected = {template["key"] for template in spec["templates"]}
    current = {item["key"]: item for item in manifest.get("templates", [])}
    if set(current) != expected:
        raise ValueError("Canonical provenance manifest does not match the rule spec template inventory.")

    with tempfile.TemporaryDirectory(prefix="canonical-template-determinism-") as temporary:
        regenerated = generate_canonical_templates(spec_path, Path(temporary) / "generated")
        regenerated_by_key = {item["key"]: item for item in regenerated["templates"]}
    results = []
    for template in spec["templates"]:
        key = template["key"]
        expected_hash = current[key]["canonicalSha256"]
        regenerated_hash = regenerated_by_key[key]["canonicalSha256"]
        current_file_hash = sha256_file(input_dir / Path(template["canonicalFile"]))
        passed = expected_hash == regenerated_hash == current_file_hash
        results.append({
            "templateKey": key,
            "canonicalSha256": current_file_hash,
            "regeneratedSha256": regenerated_hash,
            "status": "passed" if passed else "failed",
        })
    return {
        "ruleSpecVersion": spec["ruleSpecVersion"],
        "generatorVersion": GENERATOR_VERSION,
        "generatorSha256": sha256_file(Path(__file__).resolve()),
        "status": "passed" if all(item["status"] == "passed" for item in results) else "failed",
        "templates": results,
    }


def record_qa_evidence(
    output_dir: Path,
    spec_path: Path,
    structural_report_path: Path,
    deterministic_report_path: Path,
    visual_report_path: Path,
) -> dict[str, Any]:
    spec = load_spec(spec_path)
    output_dir = Path(output_dir)
    manifest_path = output_dir / "provenance.json"
    manifest = _read_json(manifest_path)
    if manifest.get("ruleSpecVersion") != spec["ruleSpecVersion"]:
        raise ValueError("Canonical manifest ruleSpecVersion does not match the QA rule spec.")
    generator_hash = sha256_file(Path(__file__).resolve())
    if manifest.get("generatorVersion") != GENERATOR_VERSION or manifest.get("generatorSha256") != generator_hash:
        raise ValueError("Canonical manifest was not generated by the current generator version and hash.")

    structural_report = _read_json(structural_report_path)
    deterministic_report = _read_json(deterministic_report_path)
    visual_report = _read_json(visual_report_path)
    for label, report in (("structural", structural_report), ("deterministic", deterministic_report), ("visual", visual_report)):
        if report.get("ruleSpecVersion") != spec["ruleSpecVersion"]:
            raise ValueError(f"{label} QA report ruleSpecVersion does not match the rule spec.")
    keys = {template["key"] for template in spec["templates"]}
    structural = _report_index(structural_report, "structural/semantic", keys)
    deterministic = _report_index(deterministic_report, "determinism", keys)
    visual = _report_index(visual_report, "visual", keys)
    manifest_templates = {item["key"]: item for item in manifest.get("templates", [])}
    if set(manifest_templates) != keys:
        raise ValueError("Canonical manifest template inventory does not match the rule spec.")

    report_paths = {
        "structural": _report_path(structural_report_path),
        "deterministic": _report_path(deterministic_report_path),
        "visual": _report_path(visual_report_path),
    }
    report_hashes = {
        "structural": sha256_file(structural_report_path),
        "deterministic": sha256_file(deterministic_report_path),
        "visual": sha256_file(visual_report_path),
    }
    renderer = visual_report.get("renderer", {})
    renderer_name = " ".join(value for value in (renderer.get("name"), renderer.get("version")) if value)
    if not renderer_name:
        raise ValueError("Visual QA report must identify the renderer and version.")

    for template in spec["templates"]:
        key = template["key"]
        record = manifest_templates[key]
        actual_hash = sha256_file(output_dir / Path(template["canonicalFile"]))
        if actual_hash != record.get("canonicalSha256"):
            raise ValueError(f"Canonical DOCX fingerprint changed since generation: {key}")
        for label, item in (("structural", structural[key]), ("deterministic", deterministic[key]), ("visual", visual[key])):
            item_status = item.get("status") or item.get(f"{label}Qa", {}).get("status")
            if item_status != "passed" or item.get("canonicalSha256") != actual_hash:
                raise ValueError(f"{label} QA evidence failed or has a mismatched canonical hash for {key}.")
        if structural[key].get("structuralQa", {}).get("status") != "passed":
            raise ValueError(f"Structural QA did not pass for {key}.")
        if structural[key].get("semanticQa", {}).get("status") != "passed":
            raise ValueError(f"Semantic QA did not pass for {key}.")
        if deterministic[key].get("regeneratedSha256") != actual_hash:
            raise ValueError(f"Determinism QA regenerated a different canonical hash for {key}.")
        if visual[key].get("visualQa", {}).get("status") != "passed":
            raise ValueError(f"Visual QA did not pass for {key}.")

        record["structuralQa"] = {"status": "passed", "reportPath": report_paths["structural"], "sha256": report_hashes["structural"]}
        record["semanticQa"] = {"status": "passed", "reportPath": report_paths["structural"], "sha256": report_hashes["structural"]}
        record["deterministicQa"] = {"status": "passed", "reportPath": report_paths["deterministic"], "sha256": report_hashes["deterministic"]}
        record["visualQa"] = {
            "status": "passed",
            "renderer": renderer_name,
            "reportPath": report_paths["visual"],
            "sha256": report_hashes["visual"],
        }

    _write_json(manifest_path, manifest)
    return manifest


def _runtime_destination(runtime_root: Path, runtime_path: str) -> Path:
    normalized = runtime_path.replace("\\", "/")
    parts = normalized.strip("/").split("/")
    if not parts or parts[0] != "templates" or any(part in ("", ".", "..") for part in parts):
        raise ValueError(f"Unsafe runtime path outside /templates/: {runtime_path}")
    root = Path(runtime_root).resolve()
    destination = root.joinpath(*parts[1:]).resolve()
    try:
        destination.relative_to(root)
    except ValueError as error:
        raise ValueError(f"Unsafe runtime path escapes the runtime template root: {runtime_path}") from error
    return destination


def sync_runtime_templates(
    output_dir: Path = DEFAULT_OUTPUT,
    spec_path: Path = DEFAULT_SPEC,
    runtime_root: Path = ROOT / "templates",
    parity_report_path: Path | None = None,
) -> dict[str, Any]:
    spec = load_spec(spec_path)
    output_dir = Path(output_dir)
    manifest_path = output_dir / "provenance.json"
    manifest = _read_json(manifest_path)
    if manifest.get("ruleSpecVersion") != spec["ruleSpecVersion"]:
        raise ValueError("Canonical manifest ruleSpecVersion does not match the runtime spec.")
    if manifest.get("generatorSha256") != sha256_file(Path(__file__).resolve()):
        raise ValueError("Canonical manifest generator hash does not match the current canonical generator.")
    records = {item["key"]: item for item in manifest.get("templates", [])}
    templates = {item["key"]: item for item in spec["templates"]}
    if set(records) != set(templates):
        raise ValueError("Canonical manifest template inventory does not match the runtime spec.")

    destinations = {}
    for key, template in templates.items():
        record = records[key]
        for qa_key in ("structuralQa", "semanticQa", "deterministicQa", "visualQa"):
            if record.get(qa_key, {}).get("status") != "passed":
                raise ValueError(f"Runtime sync requires all structural, semantic, deterministic and visual QA to pass ({key}: {qa_key}).")
        destinations[key] = _runtime_destination(runtime_root, template["runtimePath"])
        canonical = output_dir / Path(template["canonicalFile"])
        if not canonical.is_file() or sha256_file(canonical) != record.get("canonicalSha256"):
            raise ValueError(f"Canonical DOCX is missing or changed for {key}.")
        if template["runtimePath"] != record.get("runtimePath"):
            raise ValueError(f"Runtime path differs from generated provenance for {key}.")

    results = []
    for key, template in templates.items():
        record = records[key]
        canonical = output_dir / Path(template["canonicalFile"])
        destination = destinations[key]
        destination.parent.mkdir(parents=True, exist_ok=True)
        canonical_bytes = canonical.read_bytes()
        if not destination.is_file() or destination.read_bytes() != canonical_bytes:
            destination.write_bytes(canonical_bytes)
        runtime_hash = sha256_file(destination)
        if runtime_hash != record["canonicalSha256"]:
            raise ValueError(f"Runtime package is not byte-identical to canonical for {key}.")
        record["runtimeSha256"] = runtime_hash
        record["comparison"] = "byte-exact"
        record["runtimeParity"] = {"status": "passed", "comparison": "byte-exact", "runtimeSha256": runtime_hash}
        results.append({"templateKey": key, "runtimePath": template["runtimePath"], "runtimeSha256": runtime_hash, "comparison": "byte-exact", "status": "passed"})

    if parity_report_path is not None:
        parity_report = {
            "ruleSpecVersion": spec["ruleSpecVersion"],
            "status": "passed",
            "templates": results,
        }
        _write_json(parity_report_path, parity_report)
        report_record = {"reportPath": _report_path(parity_report_path), "sha256": sha256_file(parity_report_path)}
        for record in records.values():
            record["runtimeParity"].update(report_record)
    _write_json(manifest_path, manifest)
    return manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    generate = subparsers.add_parser("generate", help="generate canonical DOCX files only")
    generate.add_argument("--spec", type=Path, default=DEFAULT_SPEC)
    generate.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    determinism = subparsers.add_parser("verify-determinism", help="regenerate in a temporary directory and compare all canonical hashes")
    determinism.add_argument("--spec", type=Path, default=DEFAULT_SPEC)
    determinism.add_argument("--input-dir", type=Path, default=DEFAULT_OUTPUT)
    determinism.add_argument("--report", type=Path, default=ROOT / "qa" / "canonical-template-reports" / "deterministic.json")
    record_qa = subparsers.add_parser("record-qa", help="record passing structural, semantic, deterministic, and visual reports")
    record_qa.add_argument("--spec", type=Path, default=DEFAULT_SPEC)
    record_qa.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    record_qa.add_argument("--structural-report", type=Path, required=True)
    record_qa.add_argument("--deterministic-report", type=Path, required=True)
    record_qa.add_argument("--visual-report", type=Path, required=True)
    sync = subparsers.add_parser("sync-runtime", help="copy passed canonical packages byte-exactly into declared runtime paths")
    sync.add_argument("--spec", type=Path, default=DEFAULT_SPEC)
    sync.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    sync.add_argument("--runtime-root", type=Path, default=ROOT / "templates")
    sync.add_argument("--parity-report", type=Path, default=ROOT / "qa" / "canonical-template-reports" / "runtime-parity.json")
    args = parser.parse_args(argv)
    if args.command == "generate":
        manifest = generate_canonical_templates(args.spec, args.output_dir)
        print(f"Generated {len(manifest['templates'])} canonical DOCX files at {args.output_dir}")
        for template in manifest["templates"]:
            print(f"{template['key']} {template['canonicalSha256']} {template['canonicalFile']}")
        return 0
    if args.command == "verify-determinism":
        report = verify_determinism(args.spec, args.input_dir)
        _write_json(args.report, report)
        passed = sum(item["status"] == "passed" for item in report["templates"])
        print(f"Determinism QA: {report['status']} ({passed}/{len(report['templates'])} passed). Report: {args.report}")
        return 0 if report["status"] == "passed" else 1
    if args.command == "record-qa":
        manifest = record_qa_evidence(args.output_dir, args.spec, args.structural_report, args.deterministic_report, args.visual_report)
        print(f"Recorded structural, semantic, deterministic, and visual QA for {len(manifest['templates'])} canonical DOCX files.")
        return 0
    if args.command == "sync-runtime":
        manifest = sync_runtime_templates(args.output_dir, args.spec, args.runtime_root, args.parity_report)
        print(f"Runtime parity: {len(manifest['templates'])}/{len(manifest['templates'])} byte-exact. Parity report: {args.parity_report}")
        return 0
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
