#!/usr/bin/env python3
"""Read-only structural and semantic checks for generated canonical DOCX packages."""

from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import sys
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
GENERATOR_PATH = Path(__file__).with_name("generate-canonical-templates.py")
DEFAULT_SPEC = ROOT / "templates" / "specs" / "administrative-templates.v1.json"
DEFAULT_INPUT = ROOT / "canonical_templates" / "generated"
DEFAULT_REPORT = ROOT / "qa" / "canonical-template-reports" / "structural.json"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
PLACEHOLDER_PATTERN = re.compile(r"\{[A-ZÀ-Ỵ0-9_]+\}")
SEMANTIC_FINDING_CODES = {"missing-label", "forbidden-label", "unknown-placeholder"}


def _load_generator():
    module_spec = importlib.util.spec_from_file_location("canonical_template_generator", GENERATOR_PATH)
    if module_spec is None or module_spec.loader is None:
        raise RuntimeError(f"Cannot load generator module: {GENERATOR_PATH}")
    module = importlib.util.module_from_spec(module_spec)
    sys.modules[module_spec.name] = module
    module_spec.loader.exec_module(module)
    return module


def sha256_file(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _finding(code: str, message: str) -> dict[str, str]:
    return {"code": code, "message": message}


def _iter_cell_paragraphs(cell):
    for paragraph in cell.paragraphs:
        yield paragraph
    for table in cell.tables:
        for row in table.rows:
            for nested_cell in row.cells:
                yield from _iter_cell_paragraphs(nested_cell)


def _all_paragraphs(document):
    yield from document.paragraphs
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from _iter_cell_paragraphs(cell)
    for section in document.sections:
        for part in (section.header, section.first_page_header, section.footer, section.first_page_footer):
            yield from part.paragraphs
            for table in part.tables:
                for row in table.rows:
                    for cell in row.cells:
                        yield from _iter_cell_paragraphs(cell)


def _table_text(table) -> str:
    return "\n".join(
        paragraph.text
        for row in table.rows
        for cell in row.cells
        for paragraph in _iter_cell_paragraphs(cell)
    )


def _paragraph_style_findings(document, spec: dict[str, Any], template: dict[str, Any]) -> list[dict[str, str]]:
    common = spec["commonStyle"]
    type_schema = spec["documentTypes"][template["typeSchema"]]
    findings = []
    body_paragraphs = list(document.paragraphs)

    title_block = next((block for block in type_schema["blocks"] if block["kind"] == "title"), None)
    if title_block:
        expected_text = title_block["text"]
        title_paragraph = next((paragraph for paragraph in body_paragraphs if paragraph.text == expected_text), None)
        if title_paragraph is None:
            findings.append(_finding("missing-title", f"Document title paragraph {expected_text!r} is absent."))
        else:
            style = common["documentTitle"]
            if title_paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                findings.append(_finding("title-alignment", "Document title must be centered."))
            if not title_paragraph.runs or any(run.font.bold is not True for run in title_paragraph.runs if run.text):
                findings.append(_finding("title-bold", "Document title must be bold."))
            if any(run.font.size is None or abs(run.font.size.pt - style["fontPt"]) > 0.05 for run in title_paragraph.runs if run.text):
                findings.append(_finding("title-font-size", "Document title font size does not match the spec."))

    all_paragraphs = list(_all_paragraphs(document))
    paragraph_texts = [paragraph.text for paragraph in all_paragraphs]
    for block in type_schema["blocks"]:
        if block["kind"] != "subject":
            continue
        expected_text = block["text"]
        paragraph = next((item for item in all_paragraphs if item.text == expected_text), None)
        if paragraph is None:
            continue
        style = common[block.get("style", "subject")]
        if paragraph.alignment != WD_ALIGN_PARAGRAPH[style["alignment"].upper()]:
            findings.append(_finding("subject-alignment", f"Subject {expected_text!r} alignment does not match the spec."))
        if any(run.font.size is None or abs(run.font.size.pt - style["fontPt"]) > 0.05 for run in paragraph.runs if run.text):
            findings.append(_finding("subject-font-size", f"Subject {expected_text!r} font size does not match the spec."))
        if style.get("bold") and any(run.font.bold is not True for run in paragraph.runs if run.text):
            findings.append(_finding("subject-bold", f"Subject {expected_text!r} must be bold."))
        if style.get("italic") and any(run.font.italic is not True for run in paragraph.runs if run.text):
            findings.append(_finding("subject-italic", f"Subject {expected_text!r} must be italic."))

    national_heading = common["nationalHeading"]["text"]
    if type_schema.get("headerMode", "administrative") != "internal":
        paragraph = next((item for item in all_paragraphs if national_heading in item.text), None)
        if paragraph is None:
            findings.append(_finding("national-heading", "Administrative header lacks the national heading."))
        else:
            heading_style = common["nationalHeading"]
            if paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER:
                findings.append(_finding("national-heading-alignment", "National heading must be centered."))
            if any(run.font.bold is not True for run in paragraph.runs if run.text):
                findings.append(_finding("national-heading-bold", "National heading must be bold."))
            if any(run.font.size is None or abs(run.font.size.pt - heading_style["fontPt"]) > 0.05 for run in paragraph.runs if run.text):
                findings.append(_finding("national-heading-font-size", "National heading font size does not match the spec."))

        motto_text = common["nationalMotto"]["text"]
        motto = next((item for item in all_paragraphs if motto_text in item.text), None)
        if motto is None:
            findings.append(_finding("national-motto", "Administrative header lacks the national motto."))
        elif not _has_paragraph_bottom_rule(motto):
            findings.append(_finding("national-motto-rule", "National motto must have the specified bottom rule."))

    symbol_pattern = template.get("symbolPattern")
    if symbol_pattern:
        symbol = next((item for item in all_paragraphs if item.text == symbol_pattern), None)
        if symbol is None:
            findings.append(_finding("symbol-line", "Document symbol line is missing."))
        else:
            style = common["symbol"]
            if any(run.font.size is None or abs(run.font.size.pt - style["fontPt"]) > 0.05 for run in symbol.runs if run.text):
                findings.append(_finding("symbol-font-size", "Document symbol font size does not match the spec."))

    parent = spec["organizations"][template["organization"]]["parentAgency"]
    issuer = spec["organizations"][template["organization"]]["issuingAgency"]
    header_text = _table_text(document.tables[0]) if document.tables else ""
    if parent not in header_text or issuer not in header_text:
        findings.append(_finding("header-organization", "Header does not contain the organization hierarchy recorded in the spec."))

    issuer_paragraph = next((item for item in all_paragraphs if item.text == issuer), None)
    if issuer_paragraph is None or not _has_paragraph_bottom_rule(issuer_paragraph):
        findings.append(_finding("issuer-rule", "Issuing-agency line lacks its rule."))

    for block_kind, code, required in (
        ("recipients_signature", "recipient-signature-block", ("Nơi nhận:", "{CHUC_VU_NGUOI_KY}", "{HO_TEN_NGUOI_KY}")),
        ("meeting_signatures", "meeting-signature-block", ("CHỦ TRÌ", "THƯ KÝ")),
        ("leave_signatures", "leave-signature-block", ("NGƯỜI LÀM ĐƠN", "TRƯỞNG ĐƠN VỊ", "LÃNH ĐẠO PHÊ DUYỆT")),
        ("commitment_signatures", "commitment-signature-block", ("NGƯỜI CAM KẾT", "XÁC NHẬN CỦA ĐƠN VỊ")),
        ("copy_certification", "copy-certification-block", ("SAO Y BẢN CHÍNH", "{CHUC_VU_NGUOI_KY}")),
    ):
        if any(block["kind"] == block_kind for block in type_schema["blocks"]):
            for text in required:
                if text not in "\n".join(paragraph_texts):
                    findings.append(_finding(code, f"Required block text {text!r} is absent."))

    return findings


def _has_paragraph_bottom_rule(paragraph) -> bool:
    p_pr = paragraph._p.pPr
    if p_pr is None:
        return False
    borders = p_pr.find(qn("w:pBdr"))
    return borders is not None and borders.find(qn("w:bottom")) is not None


def _package_control_findings(path: Path) -> list[dict[str, str]]:
    findings = []
    tags: list[str] = []
    ids: list[str] = []
    with zipfile.ZipFile(path, "r") as package:
        for name in package.namelist():
            if not name.endswith(".xml"):
                continue
            try:
                root = ET.fromstring(package.read(name))
            except ET.ParseError as error:
                findings.append(_finding("invalid-xml", f"{name} is invalid XML: {error}"))
                continue
            for element in root.findall(f".//{{{W_NS}}}sdtPr/{{{W_NS}}}tag"):
                value = element.attrib.get(f"{{{W_NS}}}val")
                if value:
                    tags.append(value)
            for element in root.findall(f".//{{{W_NS}}}sdtPr/{{{W_NS}}}id"):
                value = element.attrib.get(f"{{{W_NS}}}val")
                if value:
                    ids.append(value)
            for row_height in root.findall(f".//{{{W_NS}}}trPr/{{{W_NS}}}trHeight"):
                if row_height.attrib.get(f"{{{W_NS}}}hRule") == "exact":
                    findings.append(_finding("fixed-row-height", f"{name} contains an exact fixed row height."))
    if len(tags) != len(set(tags)):
        findings.append(_finding("duplicate-control-tag", "DOCX contains duplicate content-control tags."))
    if len(ids) != len(set(ids)):
        findings.append(_finding("duplicate-control-id", "DOCX contains duplicate content-control IDs."))
    return findings


def _page_number_findings(document, spec: dict[str, Any], template: dict[str, Any]) -> list[dict[str, str]]:
    findings = []
    page_number = spec["commonStyle"]["pageNumber"]
    section = document.sections[0]
    header_xml = section.header._element.xml
    first_header_xml = section.first_page_header._element.xml
    if "fldSimple" not in header_xml or "PAGE" not in header_xml:
        findings.append(_finding("page-number-field", "Default header lacks a Word PAGE field."))
    if page_number["suppressFirstPage"] and ("fldSimple" in first_header_xml or "PAGE" in first_header_xml):
        findings.append(_finding("first-page-number", "First-page header must suppress the page number."))
    paragraph = section.header.paragraphs[0]
    expected_compensation = (section.right_margin.mm - section.left_margin.mm) / 2
    left_indent = paragraph.paragraph_format.left_indent
    right_indent = paragraph.paragraph_format.right_indent
    left_mm = left_indent.mm if left_indent is not None else 0
    right_mm = right_indent.mm if right_indent is not None else 0
    if (
        paragraph.alignment != 1
        or abs(left_mm - expected_compensation) > 0.1
        or abs(right_mm + expected_compensation) > 0.1
    ):
        findings.append(_finding("page-number-position", "Page number is not centered on the physical page width."))
    if "w:val=\"26\"" not in header_xml and "w:val='26'" not in header_xml:
        findings.append(_finding("page-number-font-size", "Page number must use the specified 13 pt size."))
    return findings


def verify_docx(path: Path, spec: dict[str, Any], template: dict[str, Any]) -> dict[str, Any]:
    path = Path(path)
    findings: list[dict[str, str]] = []
    if not path.is_file():
        return {"templateKey": template["key"], "path": str(path), "findings": [_finding("missing-file", "DOCX file does not exist.")]}
    try:
        document = Document(str(path))
    except Exception as error:  # python-docx surfaces several package/XML exceptions.
        return {"templateKey": template["key"], "path": str(path), "findings": [_finding("unopenable-docx", str(error))]}

    page = spec["commonStyle"]["page"]
    expected_margins = page["marginsMm"]
    section = document.sections[0]
    if abs(section.page_width.mm - page["widthMm"]) > 0.2 or abs(section.page_height.mm - page["heightMm"]) > 0.2:
        findings.append(_finding("page-size", "Page must be A4 with the dimensions recorded in the spec."))
    actual_margins = {
        "top": section.top_margin.mm,
        "bottom": section.bottom_margin.mm,
        "left": section.left_margin.mm,
        "right": section.right_margin.mm,
    }
    if any(abs(actual_margins[key] - expected_margins[key]) > 0.2 for key in expected_margins):
        findings.append(_finding("page-margin", f"Margins {actual_margins} do not match {expected_margins}."))

    usable_width = page["widthMm"] - expected_margins["left"] - expected_margins["right"]
    if not document.tables or len(document.tables[0].columns) != spec["commonStyle"]["headerGrid"]["columns"]:
        findings.append(_finding("header-grid", "Document body must start with a two-column organization header grid."))
    else:
        header = document.tables[0]
        width_sum = sum(column.width.mm for column in header.columns)
        if width_sum > usable_width + 0.25:
            findings.append(_finding("header-width", f"Header width {width_sum:.2f} mm exceeds usable width {usable_width:.2f} mm."))
        layout = header._tbl.tblPr.find(qn("w:tblLayout"))
        if layout is None or layout.get(qn("w:type")) != "fixed":
            findings.append(_finding("header-layout", "Header columns must use the declared fixed grid geometry."))

    allowed_placeholders = set(spec["placeholders"])
    all_text = "\n".join(paragraph.text for paragraph in _all_paragraphs(document))
    unknown_placeholders = set(PLACEHOLDER_PATTERN.findall(all_text)) - allowed_placeholders
    if unknown_placeholders:
        findings.append(_finding("unknown-placeholder", f"Undeclared placeholders: {sorted(unknown_placeholders)}."))

    type_schema = spec["documentTypes"][template["typeSchema"]]
    for label in type_schema["requiredLabels"]:
        if label not in all_text:
            findings.append(_finding("missing-label", f"Required label {label!r} is missing."))
    for label in type_schema["forbiddenLabels"]:
        if label in all_text:
            findings.append(_finding("forbidden-label", f"Forbidden label {label!r} is present."))

    font_name = spec["commonStyle"]["font"]["name"]
    for paragraph in _all_paragraphs(document):
        for run in paragraph.runs:
            if not run.text:
                continue
            if run.font.name != font_name:
                findings.append(_finding("font-name", f"Run {run.text[:40]!r} is not explicitly set to {font_name}."))
                break
            if run.font.color.rgb is None or str(run.font.color.rgb).upper() != spec["commonStyle"]["font"]["color"]:
                findings.append(_finding("font-color", f"Run {run.text[:40]!r} is not explicitly black."))
                break

    findings.extend(_paragraph_style_findings(document, spec, template))
    findings.extend(_page_number_findings(document, spec, template))
    findings.extend(_package_control_findings(path))
    semantic_findings = [finding for finding in findings if finding["code"] in SEMANTIC_FINDING_CODES]
    structural_findings = [finding for finding in findings if finding["code"] not in SEMANTIC_FINDING_CODES]
    return {
        "templateKey": template["key"],
        "canonicalFile": template["canonicalFile"],
        "canonicalSha256": sha256_file(path),
        "structuralQa": {"status": "passed" if not structural_findings else "failed"},
        "semanticQa": {"status": "passed" if not semantic_findings else "failed"},
        "status": "passed" if not findings else "failed",
        "findings": findings,
    }


def verify_all(spec_path: Path = DEFAULT_SPEC, input_dir: Path = DEFAULT_INPUT) -> dict[str, Any]:
    generator = _load_generator()
    spec = generator.load_spec(spec_path)
    results = [
        verify_docx(input_dir / Path(template["canonicalFile"]), spec, template)
        for template in spec["templates"]
    ]
    return {
        "ruleSpecVersion": spec["ruleSpecVersion"],
        "status": "passed" if all(result["status"] == "passed" for result in results) else "failed",
        "templates": results,
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--all", action="store_true", help="verify every spec-declared canonical document")
    parser.add_argument("--spec", type=Path, default=DEFAULT_SPEC)
    parser.add_argument("--input-dir", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    args = parser.parse_args(argv)
    if not args.all:
        parser.error("Pass --all to run the full canonical DOCX QA suite.")
    report = verify_all(args.spec, args.input_dir)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, sort_keys=True, indent=2) + "\n", encoding="utf-8", newline="\n")
    failed = sum(result["status"] != "passed" for result in report["templates"])
    print(f"Structural and semantic QA: {report['status']} ({len(report['templates']) - failed}/{len(report['templates'])} passed). Report: {args.report}")
    return 0 if report["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
