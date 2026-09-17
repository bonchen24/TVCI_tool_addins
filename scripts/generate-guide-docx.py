import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout_box(doc, text, title="LƯU Ý QUAN TRỌNG:", hex_bg="F0F7FF", hex_border="0284C7"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, hex_bg)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=180)
    
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="single" w:sz="36" w:space="0" w:color="{hex_border}"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
    cell._tc.get_or_add_tcPr().append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    run_title = p.add_run(f"📌 {title} ")
    run_title.bold = True
    run_title.font.name = "Times New Roman"
    run_title.font.size = Pt(11.5)
    run_title.font.color.rgb = RGBColor(2, 132, 199)
    
    run_text = p.add_run(text)
    run_text.font.name = "Times New Roman"
    run_text.font.size = Pt(11)
    run_text.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph()

def create_guide_document(output_path):
    doc = Document()
    
    # Page setup: A4
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.79) # 2cm
        section.bottom_margin = Inches(0.79) # 2cm
        section.left_margin = Inches(0.98) # 2.5cm
        section.right_margin = Inches(0.79) # 2cm
        
        # Header & Footer
        footer = section.footer
        p_ft = footer.paragraphs[0]
        p_ft.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        f_run = p_ft.add_run("Tài liệu hướng dẫn TVCI Word Tools v1.0 | ")
        f_run.font.name = "Times New Roman"
        f_run.font.size = Pt(9)
        f_run.font.italic = True
        f_run.font.color.rgb = RGBColor(120, 120, 120)

    # Styles default
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(30, 41, 59)

    # --- COVER / HEADER BANNER ---
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_after = Pt(2)
    r_sub = p_top.add_run("CÔNG TY CỔ PHẦN TVCI • VIỆN IEMM\nTRUNG TÂM PHÁT TRIỂN CÔNG NGHỆ VÀ CHUYỂN ĐỔI SỐ")
    r_sub.font.name = "Times New Roman"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(71, 85, 105)
    
    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_after = Pt(18)
    r_div = p_div.add_run("─────────── • ───────────")
    r_div.font.color.rgb = RGBColor(148, 163, 184)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(6)
    p_title.paragraph_format.space_after = Pt(8)
    r_title = p_title.add_run("TÀI LIỆU HƯỚNG DẪN CÀI ĐẶT & SỬ DỤNG\nBỘ CÔNG CỤ TVCI WORD TOOLS")
    r_title.font.name = "Times New Roman"
    r_title.font.size = Pt(20)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    r_sub2 = p_sub.add_run("Tự động hóa chuẩn hóa văn bản hành chính theo Nghị định 30/2020/NĐ-CP, thể thức Đảng theo Hướng dẫn 36-HD/VPTW và Trợ lý AI Soạn thảo thông minh")
    r_sub2.font.name = "Times New Roman"
    r_sub2.font.size = Pt(12)
    r_sub2.font.italic = True
    r_sub2.font.color.rgb = RGBColor(71, 85, 105)

    # Info table
    tbl_info = doc.add_table(rows=4, cols=2)
    tbl_info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("Phiên bản ứng dụng:", "TVCI Word Add-in v1.0.0 (Bản phân phối Doanh nghiệp)"),
        ("Môi trường tương thích:", "Microsoft Word 2016 / 2019 / 2021 / Microsoft 365 (Windows 10/11)"),
        ("Chế độ hoạt động:", "Độc lập Cục bộ (Offline Standalone) & Hỗ trợ Web Cloud"),
        ("Ngày phát hành:", "Năm 2026")
    ]
    for idx, (lbl, val) in enumerate(info_data):
        row = tbl_info.rows[idx]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        set_cell_background(c0, "F8FAFC")
        set_cell_background(c1, "FFFFFF")
        set_cell_margins(c0, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        
        p0 = c0.paragraphs[0]
        p0.paragraph_format.space_after = Pt(0)
        r0 = p0.add_run(lbl)
        r0.font.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = RGBColor(51, 65, 85)
        
        p1 = c1.paragraphs[0]
        p1.paragraph_format.space_after = Pt(0)
        r1 = p1.add_run(val)
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(15, 23, 42)

    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    
    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(6)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(14, 116, 144)
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 58, 138)
        return h

    def add_heading_3(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(2)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.italic = True
        run.font.color.rgb = RGBColor(51, 65, 85)
        return h

    def add_p(text, bold_prefix=None, italic=False, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.2
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = "Times New Roman"
            r_b.font.bold = True
            r_b.font.size = Pt(12)
            r_b.font.color.rgb = RGBColor(15, 23, 42)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.italic = italic
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.font.name = "Times New Roman"
            r_b.font.bold = True
            r_b.font.size = Pt(12)
            r_b.font.color.rgb = RGBColor(15, 23, 42)
        r = p.add_run(text)
        r.font.name = "Times New Roman"
        r.font.size = Pt(12)
        r.font.color.rgb = RGBColor(30, 41, 59)
        return p

    # --- PHẦN I ---
    add_heading_1("I. GIỚI THIỆU CHUNG VỀ BỘ CÔNG CỤ TVCI WORD TOOLS")
    add_p("TVCI Word Tools là tiện ích mở rộng (Add-in) được phát triển dành riêng cho Microsoft Word, nhằm tự động hóa quy trình soạn thảo văn bản, giảm thiểu 90% lỗi định dạng hình thức và nâng cao năng suất làm việc của cán bộ, chuyên viên, nhân viên văn phòng.")
    
    add_p("Các phân hệ chức năng cốt lõi của bộ công cụ gồm có:")
    add_bullet("Cung cấp đầy đủ các mẫu văn bản hành chính theo Nghị định số 30/2020/NĐ-CP (Công văn, Thông báo, Tờ trình, Quyết định, Báo cáo, Kế hoạch, Giấy mời...). Tự động chuẩn hóa lề trang A4 (Trên 20-25mm, Dưới 20-25mm, Trái 30-35mm, Phải 15-20mm), font chữ Times New Roman, dãn dòng 1.3 - 1.4.", "1. Thư viện Biểu mẫu Hành chính Nhà nước: ")
    add_bullet("Tích hợp hệ thống mẫu văn bản của các cấp ủy Đảng theo quy định Hướng dẫn số 36-HD/VPTW, đảm bảo chính xác vị trí, cỡ chữ, căn lề và dấu hiệu tiêu đề thể thức Đảng.", "2. Thư viện Biểu mẫu Thể thức Đảng: ")
    add_bullet("Hệ thống biểu mẫu đặc thù của Viện IEMM và Công ty TVCI (Công văn đối tác, Thông báo nội bộ, Báo cáo kỹ thuật, Đơn xin nghỉ phép, Thư mời...), đảm bảo nhận diện thương hiệu thống nhất.", "3. Hệ thống Biểu mẫu Viện IEMM & Doanh nghiệp: ")
    add_bullet("Trợ lý trí tuệ nhân tạo (hỗ trợ OpenAI, Gemini, Claude, DeepSeek, Local LLM) giúp dự thảo nội dung văn bản từ yêu cầu ngắn, tự động phân tích cấu trúc đoạn và tự động điền trực tiếp vào biểu mẫu Word đang chọn chỉ bằng 1 cú nhấp chuột.", "4. AI Workspace (Trợ lý Soạn thảo thông minh): ")
    add_bullet("Kiểm tra và sửa lỗi chính tả tiếng Việt tự động, chuẩn hóa khoảng trắng kép, sửa dấu ngắt câu theo ngữ pháp tiếng Việt, chuyển đổi hoa/thường, chuyển mã font.", "5. Bộ tiện ích Soát lỗi & Chuẩn hóa: ")

    add_callout_box(
        doc,
        "Bộ công cụ TVCI Word Tools hoạt động trực tiếp trong Microsoft Word thông qua nền tảng hiện đại Office Web Add-ins (sử dụng Edge WebView2). Tiện ích có khả năng chạy độc lập CỤC BỘ (Offline) trên từng máy tính hoặc triển khai tập trung qua máy chủ nội bộ Intranet/Internet của đơn vị.",
        title="ĐẶC ĐIỂM CÔNG NGHỆ"
    )

    # --- PHẦN II ---
    add_heading_1("II. HƯỚNG DẪN QUẢN LÝ VÒNG ĐỜI ỨNG DỤNG: CÀI ĐẶT, SỬA CHỮA, CẬP NHẬT VÀ GỠ BỎ")
    add_p("Bộ công cụ TVCI Word Tools được đóng gói theo chuẩn phần mềm doanh nghiệp, hỗ trợ đầy đủ 4 giai đoạn trong vòng đời ứng dụng trên máy tính của người dùng:")

    add_heading_2("1. Quy trình Cài đặt mới (Installation) bằng file TVCI_Word_Addin_Setup.exe")
    add_p("Áp dụng khi cài đặt lần đầu tiên trên bất kỳ máy tính Windows nào (không yêu cầu cài đặt Node.js hay phần mềm lập trình):")
    add_bullet("Sao chép file TVCI_Word_Addin_Setup_v1.0.0.exe sang máy tính đích (thông qua USB, thư mục mạng chia sẻ LAN hoặc tải từ kho nội bộ).", "Bước 1 - Chuẩn bị: ")
    add_bullet("Đóng Microsoft Word (nếu đang mở). Nhấp đúp vào file cài đặt. Trình cài đặt hỗ trợ cài cho người dùng hiện tại (không cần quyền Administrator) hoặc cài đặt toàn hệ thống.", "Bước 2 - Khởi chạy: ")
    add_bullet("Giao diện cài đặt hiện lên: Nhấn 'Next', chọn thư mục cài đặt (mặc định tại %LOCALAPPDATA%\\Programs\\TVCI Word Tools hoặc Program Files).", "Bước 3 - Chọn vị trí: ")
    add_bullet("Tại mục 'Tùy chọn bổ sung', bạn có thể tích chọn 'Tự động khởi động dịch vụ ngầm cùng Windows' để tiện ích luôn sẵn sàng mỗi khi bật máy.", "Bước 4 - Khởi động cùng Windows: ")
    add_bullet("Nhấn 'Install' (Cài đặt). Hệ thống sẽ tự động sao chép các file biểu mẫu, biên dịch máy chủ tĩnh HTTPS tvci-host.exe (~10 KB), tự động cấp chứng chỉ bảo mật SSL cho localhost và đăng ký Add-in vào Word trong 3-5 giây.", "Bước 5 - Tiến trình tự động: ")
    add_bullet("Nhấn 'Finish' (Hoàn tất). Mở Microsoft Word lên: Tab TVCI Tools sẽ hiển thị ngay trên thanh Ribbon sẵn sàng làm việc.", "Bước 6 - Xác nhận: ")

    add_heading_2("2. Quy trình Sửa chữa (Repair) khi Add-in bị lỗi hoặc mất Ribbon")
    add_p("Trong quá trình sử dụng, nếu Word gặp sự cố crash bất thường, mất nút TVCI Tools trên thanh Ribbon, hoặc Taskpane báo lỗi màn hình trắng 'Không thể kết nối máy chủ', bạn hãy sử dụng tính năng Sửa chữa (Repair):")
    add_bullet("Vào Start Menu -> Tìm thư mục TVCI Word Tools -> Nhấp chọn 'Sửa chữa Add-in (Repair)'. Hoặc vào thư mục cài đặt và nhấp đúp file repair.cmd.", "Cách 1 (Nhanh nhất): ")
    add_bullet("Vào Windows Settings -> Apps -> Installed Apps (hoặc Control Panel > Programs and Features) -> Tìm 'TVCI Word Tools' -> Bấm vào dấu ba chấm (...) và chọn 'Modify / Repair'.", "Cách 2 (Qua Control Panel): ")
    add_bullet("Tự động quét và tắt các tiến trình máy chủ bị treo; Tự động cài lại chứng chỉ SSL Localhost vào kho tin cậy; Ghi đè lại Registry nhận diện Add-in của Word; Tự động dọn sạch bộ nhớ cache bị lỗi của Office (%LOCALAPPDATA%\\Microsoft\\Office\\16.0\\Wef); Khởi động lại dịch vụ ngầm sạch sẽ.", "Hành động hệ thống thực hiện: ")
    add_p("Sau khi màn hình thông báo 'SỬA CHỮA HOÀN TẤT THÀNH CÔNG', bạn chỉ việc mở lại Word là Add-in hoạt động bình thường.")

    add_heading_2("3. Quy trình Cập nhật phiên bản mới (Update / Upgrade)")
    add_p("Khi Ban Phát triển TVCI phát hành phiên bản mới (ví dụ bổ sung mẫu văn bản mới, nâng cấp trí tuệ nhân tạo AI hoặc sửa lỗi):")
    add_bullet("Tải file cài đặt phiên bản mới (ví dụ: TVCI_Word_Addin_Setup_v1.1.0.exe).", "Bước 1: ")
    add_bullet("Chạy trực tiếp file cài đặt mới đè lên phiên bản cũ mà KHÔNG CẦN gỡ bản cũ trước.", "Bước 2: ")
    add_bullet("Hệ thống tự động phát hiện phiên bản đã cài đặt: Tự động dừng an toàn dịch vụ cũ, cập nhật toàn bộ thư mục giao diện, biểu mẫu và manifest mới.", "Bước 3 - Nâng cấp tự động: ")
    add_bullet("BẢO TOÀN DỮ LIỆU CÁ NHÂN: Toàn bộ cấu hình khóa API (OpenAI, Gemini, DeepSeek...) và lịch sử làm việc được lưu trong vùng dữ liệu an toàn của người dùng sẽ ĐƯỢC GIỮ NGUYÊN 100%, không bị mất sau khi cập nhật.", "Bước 4 - Bảo mật thiết lập: ")

    add_heading_2("4. Quy trình Gỡ cài đặt hoàn toàn (Uninstallation)")
    add_p("Khi không còn nhu cầu sử dụng hoặc cần chuyển đổi máy:")
    add_bullet("Vào Windows Settings -> Apps -> Installed apps -> Tìm 'TVCI Word Tools' -> Nhấn 'Uninstall'. Hoặc chạy file uninstall.cmd trong thư mục cài đặt.", "Thực hiện: ")
    add_bullet("Trình gỡ cài đặt sẽ tự động: Dừng và tắt máy chủ ngầm tvci-host.exe; Xóa bỏ đăng ký Add-in trong Microsoft Word để thanh Ribbon sạch sẽ hoàn toàn; Xóa bỏ mục tự khởi động trong Windows Startup; Xóa toàn bộ file và bộ nhớ đệm Office WEF.", "Kết quả: ")

    # --- PHẦN III ---
    add_heading_1("III. HƯỚNG DẪN CẤU HÌNH VÀ SỬ DỤNG TRỢ LÝ AI SOẠN THẢO (AI WORKSPACE)")
    add_p("AI Workspace là tính năng đột phá giúp văn thư và chuyên viên tự động hóa việc dự thảo nội dung văn bản phức tạp và phân tích tự động điền vào đúng biểu mẫu hành chính.")

    add_heading_2("1. Thiết lập Cấu hình AI (API Key & Nhà cung cấp)")
    add_p("Để sử dụng trí tuệ nhân tạo, bạn chỉ cần cấu hình khóa API một lần duy nhất. Các lần sau tiện ích sẽ tự động lưu trữ an toàn trên máy tính của bạn:")
    add_bullet("Mở Taskpane TVCI Tools -> Nhấp vào biểu tượng ⚙️ (Cài đặt) hoặc mục Cấu hình AI.", "Bước 1: ")
    add_bullet("Lựa chọn Nhà cung cấp AI mong muốn:", "Bước 2: ")
    
    # Provider table
    tbl_ai = doc.add_table(rows=5, cols=3)
    tbl_ai.alignment = WD_TABLE_ALIGNMENT.CENTER
    ai_headers = ["Nhà cung cấp", "Các dòng Model đề xuất", "Mục đích sử dụng tối ưu"]
    for i, h in enumerate(ai_headers):
        c = tbl_ai.rows[0].cells[i]
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=80, bottom=80, left=100, right=100)
        p = c.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.font.bold = True
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(255, 255, 255)

    ai_data = [
        ("Google Gemini (Khuyên dùng)", "gemini-1.5-flash, gemini-2.0-flash", "Tốc độ phản hồi cực nhanh, miễn phí hạn mức lớn, ngữ cảnh tiếng Việt chuẩn xác."),
        ("OpenAI (ChatGPT)", "gpt-4o, gpt-4o-mini", "Khả năng lập luận sâu, cấu trúc câu văn hành chính chặt chẽ, trang trọng."),
        ("DeepSeek", "deepseek-chat, deepseek-reasoner", "Chi phí cực thấp, tư duy logic và diễn đạt tiếng Việt phong phú."),
        ("Local LLM (Ollama / vLLM)", "Llama-3, Qwen-2.5, PhoGPT", "Chạy nội bộ 100% trên máy chủ đơn vị, bảo mật tuyệt đối cho cơ quan nhà nước.")
    ]
    for row_idx, row_content in enumerate(ai_data, start=1):
        for col_idx, text in enumerate(row_content):
            c = tbl_ai.rows[row_idx].cells[col_idx]
            set_cell_background(c, "F8FAFC" if row_idx % 2 == 1 else "FFFFFF")
            set_cell_margins(c, top=60, bottom=60, left=80, right=80)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(text)
            r.font.size = Pt(10.5)
            r.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    add_bullet("Dán mã API Key của bạn vào ô 'API Key' và nhấn 'Kiểm tra & Lưu cấu hình'. Khi hệ thống báo tích xanh thành công là bạn đã sẵn sàng.", "Bước 3: ")

    add_heading_2("2. Quy trình Soạn thảo và Tự động áp dụng vào Biểu mẫu")
    add_p("Giao diện AI Workspace hiện đại được thiết kế theo luồng tương tác đơn dòng (Unified Chat Stream) gọn gàng, tinh tế và tối ưu không gian màn hình:")
    
    add_bullet("Tại ô nhập lệnh (Textarea), gõ yêu cầu của bạn. Ví dụ: 'Soạn giúp tôi một Thông báo nghỉ lễ Quốc khánh 02/9 của Viện IEMM cho toàn thể cán bộ nhân viên, nghỉ từ 31/8 đến hết 03/9'. Bạn cũng có thể bấm nút 📎 Ghim ngữ cảnh để đưa đoạn văn bản đang bôi đen trên Word vào làm đầu vào cho AI.", "Bước 1 - Ra lệnh: ")
    add_bullet("Nhấn nút ⚡ GỬI AI SOẠN THẢO (hoặc phím tắt Ctrl + Enter). AI sẽ ngay lập tức phân tích yêu cầu, tư duy thể thức và đưa ra nội dung dự thảo hoàn chỉnh trong khung chat.", "Bước 2 - AI xử lý: ")
    add_bullet("AI sẽ tự động đối sánh nội dung với thư viện biểu mẫu để đề xuất mẫu văn bản phù hợp nhất (Ví dụ: Thông báo IEMM hoặc Công văn TVCI). Bạn có thể bấm chọn đổi sang mẫu khác từ danh sách thả xuống nếu muốn.", "Bước 3 - Nhận diện mẫu: ")
    add_bullet("Dưới câu trả lời của AI có sẵn các nút hành động tiện lợi:", "Bước 4 - Thao tác kết quả: ")

    add_bullet("ĐÂY LÀ TÍNH NĂNG MẠNH MẼ NHẤT! AI sẽ tự động bóc tách các trường: Số ký hiệu, Trích yếu, Kính gửi, Căn cứ pháp lý, Nội dung các điều khoản, Nơi nhận... và điền đồng bộ vào toàn bộ biểu mẫu văn bản trên trang Word hiện tại.", "[🚀 Áp dụng vào biểu mẫu]: ")
    add_bullet("Mở cửa sổ Form điền chi tiết để bạn rà soát hoặc sửa đổi từng trường dữ liệu trước khi chèn vào Word.", "[📝 Mở Form]: ")
    add_bullet("Thay thế nội dung đoạn văn bản bạn đang bôi đen trên trang Word bằng câu trả lời của AI.", "[📥 Thay đoạn chọn]: ")
    add_bullet("Chèn toàn bộ nội dung vừa soạn thảo ngay tại vị trí con trỏ chuột đang đứng trong tài liệu.", "[➕ Chèn con trỏ]: ")
    add_bullet("Yêu cầu AI viết lại theo phong cách trang trọng hơn, ngắn gọn hơn hoặc bổ sung thêm căn cứ pháp lý.", "[🔄 Soạn lại]: ")

    # --- PHẦN IV ---
    add_heading_1("IV. HƯỚNG DẪN SỬ DỤNG THƯ VIỆN BIỂU MẪU & CÔNG CỤ SOÁT LỖI")
    
    add_heading_2("1. Thư viện Biểu mẫu chuẩn (Hành chính Nhà nước, Đảng & Doanh nghiệp)")
    add_p("Khi cần bắt đầu một văn bản mới hoàn toàn chuẩn quy định mà không cần nhớ vị trí từng thành phần:")
    add_bullet("Chuyển sang tab 'Biểu mẫu' trên Taskpane TVCI Tools.", "Bước 1: ")
    add_bullet("Lựa chọn nhóm danh mục: 'Hành chính Nhà nước (NĐ 30)', 'Thể thức Đảng', 'Biểu mẫu Viện IEMM' hoặc 'Biểu mẫu Công ty TVCI'.", "Bước 2: ")
    add_bullet("Chọn loại văn bản mong muốn (ví dụ: Quyết định cá biệt, Công văn, Tờ trình, Báo cáo tình hình...).", "Bước 3: ")
    add_bullet("Nhấn 'Chèn biểu mẫu vào văn bản'. Toàn bộ cấu trúc gồm Quốc hiệu, Tiêu ngữ, Tên cơ quan, Số ký hiệu, Bảng nội dung, Chữ ký và Nơi nhận sẽ được chèn chuẩn xác 100% đến từng milimet lề trang và kích thước font chữ.", "Bước 4: ")

    add_heading_2("2. Công cụ Căn chỉnh thể thức & Soát lỗi chính tả")
    add_bullet("Nếu bạn có một văn bản cũ hoặc văn bản do đối tác gửi bị lệch lề, sai cỡ chữ, dãn dòng lộn xộn: Chỉ cần mở văn bản trong Word và nhấn nút 'Chuẩn hóa định dạng' trên Ribbon. Bộ công cụ sẽ tự động đặt lại lề A4, đổi toàn bộ về Times New Roman, cỡ 13-14pt, thụt dòng chuẩn 1cm và căn đều hai bên (Justify).", "Chuẩn hóa định dạng tức thì: ")
    add_bullet("Quét toàn bộ tài liệu hoặc vùng chọn để phát hiện các lỗi dính chữ, gõ nhầm telex (như 'khôg', 'đc', 'trường hơp'), khoảng trắng trước dấu chấm/phẩy, dấu ngoặc không đóng... và gợi ý sửa chỉ bằng một nút bấm.", "Soát lỗi chính tả & Ngữ pháp: ")

    # --- PHẦN V ---
    add_heading_1("V. XỬ LÝ SỰ CỐ THƯỜNG GẶP (TROUBLESHOOTING)")
    
    faq_data = [
        ("1. Không thấy Tab 'TVCI Tools' xuất hiện trên Ribbon của Word", 
         "Nguyên nhân: Word chưa kịp nạp Registry hoặc đang mở trong lúc cài đặt.\n"
         "Cách xử lý:\n"
         "- Đóng hoàn toàn Microsoft Word (kiểm tra trong Task Manager đảm bảo không còn tiến trình WINWORD.EXE) rồi mở lại.\n"
         "- Nếu vẫn chưa thấy, vào menu Word: File > Options > Add-ins > Tại mục Manage chọn 'COM Add-ins' hoặc kiểm tra 'Office Add-ins'.\n"
         "- Chạy lại file setup.cmd (Run as administrator) để đăng ký lại."),
        
        ("2. Taskpane báo màn hình trắng hoặc 'Không thể kết nối đến máy chủ localhost:38473'",
         "Nguyên nhân: Dịch vụ chạy ngầm tvci-host.exe chưa được khởi động.\n"
         "Cách xử lý:\n"
         "- Nhấn tổ hợp phím Ctrl + Shift + Esc mở Task Manager, kiểm tra xem có tiến trình tvci-host.exe (hoặc node.exe) đang chạy hay không.\n"
         "- Mở thư mục cài đặt C:\\Program Files\\TVCI Word Tools và nhấp đúp vào tvci-host.exe để khởi động lại dịch vụ.\n"
         "- Hoặc nhấp đúp file repair.cmd trong thư mục cài đặt để hệ thống tự động sửa chữa."),
        
        ("3. Cảnh báo lỗi bảo mật chứng chỉ SSL khi mở Add-in",
         "Nguyên nhân: Windows chưa đưa chứng chỉ SSL tự ký của localhost vào danh sách tin cậy.\n"
         "Cách xử lý: Nhấp chuột phải vào file chứng chỉ tvci-cert.cer trong thư mục cài đặt -> Chọn 'Install Certificate' -> Chọn 'Local Machine' -> Chọn đặt chứng chỉ vào thư mục 'Trusted Root Certification Authorities' (Cơ quan cấp chứng chỉ gốc đáng tin cậy) -> Nhấn Hoàn tất."),
         
        ("4. Muốn gỡ bỏ hoàn toàn bộ công cụ khỏi máy tính",
         "Cách xử lý:\n"
         "- Vào Windows Settings > Apps > Installed apps (hoặc Control Panel > Programs and Features).\n"
         "- Tìm mục 'TVCI Word Tools' và nhấn 'Uninstall'.\n"
         "- Trình gỡ cài đặt sẽ tự động hủy đăng ký khỏi Microsoft Word, dừng dịch vụ chạy ngầm và dọn dẹp sạch sẽ hệ thống.")
    ]

    for q, a in faq_data:
        add_heading_3(q)
        add_p(a, space_after=8)

    # --- FOOTER SIGNATURE ---
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_s1 = p_sign.add_run("BAN BIÊN SOẠN & PHÁT TRIỂN CÔNG NGHỆ TVCI\n")
    r_s1.font.bold = True
    r_s1.font.size = Pt(12)
    r_s2 = p_sign.add_run("Email hỗ trợ: support@tvci.vn • Hotline: 1900.xxxx\nWebsite: https://tvci.vn")
    r_s2.font.italic = True
    r_s2.font.size = Pt(11)
    r_s2.font.color.rgb = RGBColor(100, 116, 139)

    doc.save(output_path)
    print(f"Document created successfully at: {output_path}")

if __name__ == "__main__":
    target = os.path.join(r"e:\CODING\TVCI_word_addins", "HUONG_DAN_CAI_DAT_VA_SU_DUNG_TVCI_WORD_ADDIN.docx")
    create_guide_document(target)
